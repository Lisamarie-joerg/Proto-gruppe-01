"""Etappe 4: Excel-Analyse (pro Stab + Summen) aus _etappe4_ergebnisse.json."""
import json
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

D = json.load(open('Abgabe_23.10/_etappe4_ergebnisse.json', encoding='utf-8'))
bars = D['bars']
sections = D['sections']
groups = D['groups']
proc = D['processes']
total = D['total_co2e_kg']

wb = openpyxl.Workbook()
hdr_fill = PatternFill('solid', fgColor='1F4E79')
hdr_font = Font(color='FFFFFF', bold=True)
title_font = Font(bold=True, size=13)
thin = Border(*[Side(style='thin', color='B0B0B0')] * 4)

def style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.fill = hdr_fill
        cell.font = hdr_font
        cell.alignment = Alignment(horizontal='center', vertical='center',
                                   wrap_text=True)
        cell.border = thin

# ---------- Sheet 1: Pro Stab ----------
ws = wb.active
ws.title = 'Pro Stab (A1-A3)'
ws['A1'] = 'Embodied Carbon – CO₂e je Stab, Modul A1-A3 (OEKOBAU.DAT, live)'
ws['A1'].font = title_font
ws['A2'] = ('Pavillon Stockerau | 736 Stäbe | Quellen: RH7 Stockerau RFEM-Export '
            '(Achslängen) + OEKOBAU.DAT GWP-total')
ws['A2'].font = Font(italic=True, size=9)

cols = ['Nr.', 'Stab-ID (GUID)', 'Typ', 'Querschnitt / Profil',
        'Länge (m)', 'Querschnittsfläche A (mm²)', 'Volumen V (m³)',
        'Masse Stahl (kg)', 'Betonvolumen (m³)', 'Betonmasse (kg)',
        'Bewehrung Betonstahl (kg)',
        'CO₂e Stahl (kg)', 'CO₂e Beton (kg)', 'CO₂e Bewehrung (kg)',
        'CO₂e A1-A3 gesamt (kg)']
HR = 4
for j, h in enumerate(cols, 1):
    ws.cell(row=HR, column=j, value=h)
style_header(ws, HR, len(cols))

for i, r in enumerate(sorted(bars, key=lambda x: (x['section'], x['id'])), 1):
    row = HR + i
    vals = [i, r['id'], r['typ'], r['section'], r['L_m'], r['A_mm2'],
            round(r['V_m3'], 6),
            round(r['m_stahl_kg'], 3), round(r['V_beton_m3'], 6),
            round(r['m_beton_kg'], 2), round(r['m_bewehrung_kg'], 2),
            round(r['co2e_stahl'], 4), round(r['co2e_beton'], 4),
            round(r['co2e_bewehrung'], 4), round(r['co2e_total_kg'], 4)]
    for j, v in enumerate(vals, 1):
        cell = ws.cell(row=row, column=j, value=v)
        cell.border = thin
        if j >= 5:
            cell.number_format = '#,##0.000' if j in (7, 9) else '#,##0.00'
srow = HR + len(bars) + 1
ws.cell(row=srow, column=4, value='SUMME (736 Stäbe)').font = Font(bold=True)
sums = {5: sum(r['L_m'] for r in bars),
        7: sum(r['V_m3'] for r in bars),
        8: sum(r['m_stahl_kg'] for r in bars),
        9: sum(r['V_beton_m3'] for r in bars),
        10: sum(r['m_beton_kg'] for r in bars),
        11: sum(r['m_bewehrung_kg'] for r in bars),
        12: sum(r['co2e_stahl'] for r in bars),
        13: sum(r['co2e_beton'] for r in bars),
        14: sum(r['co2e_bewehrung'] for r in bars),
        15: total}
for j, v in sums.items():
    cell = ws.cell(row=srow, column=j, value=round(v, 4))
    cell.font = Font(bold=True)
    cell.number_format = '#,##0.00'
    cell.border = thin
ws.freeze_panes = 'A5'
widths = [6, 38, 22, 26, 10, 14, 12, 13, 13, 13, 15, 12, 12, 14, 14]
for j, w in enumerate(widths, 1):
    ws.column_dimensions[get_column_letter(j)].width = w

# ---------- Sheet 2: Querschnitte ----------
ws2 = wb.create_sheet('Querschnitte')
ws2['A1'] = 'Ergebnis je Querschnitt (A1-A3)'
ws2['A1'].font = title_font
cols2 = ['Nr.', 'Typ', 'Querschnitt / Profil', 'Anzahl Stäbe',
         'Gesamtlänge (m)', 'Volumen (m³)', 'Masse Stahl (kg)',
         'Betonmasse (kg)', 'Bewehrung (kg)', 'CO₂e A1-A3 (kg)']
for j, h in enumerate(cols2, 1):
    ws2.cell(row=3, column=j, value=h)
style_header(ws2, 3, len(cols2))
for i, (sec, s) in enumerate(sorted(sections.items()), 1):
    typ = next(r['typ'] for r in bars if r['section'] == sec)
    vals = [i, typ, sec, s['n'], round(s['L'], 3), round(s['V'], 5),
            round(s['m_stahl'], 2), round(s['m_beton'], 2),
            round(s['m_bew'], 2), round(s['co2e'], 2)]
    for j, v in enumerate(vals, 1):
        cell = ws2.cell(row=3 + i, column=j, value=v)
        cell.border = thin
        if j >= 5:
            cell.number_format = '#,##0.000' if j in (5, 6) else '#,##0.00'
for j, w in enumerate([6, 22, 28, 12, 14, 12, 15, 14, 13, 15], 1):
    ws2.column_dimensions[get_column_letter(j)].width = w

# ---------- Sheet 3: Materialgruppen ----------
ws3 = wb.create_sheet('Materialgruppen')
ws3['A1'] = 'Ergebnis je Materialgruppe – Embodied Carbon A1-A3'
ws3['A1'].font = title_font
cols3 = ['Materialgruppe', 'OEKOBAU.DAT-Datensatz (Standard)',
         'Anzahl Stäbe', 'Volumen (m³)', 'Masse Stahl (kg)',
         'Betonmasse (kg)', 'Bewehrung Betonstahl (kg)',
         'CO₂e A1-A3 (kg)', 'Anteil (%)']
for j, h in enumerate(cols3, 1):
    ws3.cell(row=3, column=j, value=h)
style_header(ws3, 3, len(cols3))
DS_NAME = {'Baustahl': 'Baustahl: Profile und Stabstahl',
           'Beton': 'Beton C25/30 + Betonstahl in Stäben (Bewehrung)'}
for i, (gname, g) in enumerate(sorted(groups.items()), 1):
    vals = [gname, DS_NAME[gname], g['n'], round(g['V'], 4),
            round(g['m_stahl'], 1), round(g['m_beton'], 1),
            round(g['m_bew'], 1), round(g['co2e'], 2),
            round(100 * g['co2e'] / total, 1)]
    for j, v in enumerate(vals, 1):
        cell = ws3.cell(row=3 + i, column=j, value=v)
        cell.border = thin
        if j >= 4:
            cell.number_format = '#,##0.0' if j < 8 else '#,##0.00'
r = 3 + len(groups) + 1
ws3.cell(row=r, column=1, value='GESAMT').font = Font(bold=True)
for j, v in [(3, sum(g['n'] for g in groups.values())),
             (4, round(sum(g['V'] for g in groups.values()), 4)),
             (5, round(sum(g['m_stahl'] for g in groups.values()), 1)),
             (6, round(sum(g['m_beton'] for g in groups.values()), 1)),
             (7, round(sum(g['m_bew'] for g in groups.values()), 1)),
             (8, round(total, 2)), (9, 100.0)]:
    cell = ws3.cell(row=r, column=j, value=v)
    cell.font = Font(bold=True)
    cell.number_format = '#,##0.00' if j == 8 else '#,##0.0'
    cell.border = thin
ws3.cell(row=r + 2, column=1,
         value=f'Gesamtes Embodied Carbon (A1-A3): {total:,.2f} kg CO₂e '
               f'= {total/1000:,.3f} t CO₂e').font = Font(bold=True, size=12)
for j, w in enumerate([14, 42, 12, 12, 15, 14, 20, 15, 10], 1):
    ws3.column_dimensions[get_column_letter(j)].width = w

# ---------- Sheet 4: OEKOBAU.DAT ----------
ws4 = wb.create_sheet('OEKOBAU.DAT')
ws4['A1'] = 'OEKOBAU.DAT-Datensätze (live abgefragt)'
ws4['A1'].font = title_font
ws4['A2'] = ('Schnittstelle: https://oekobaudat.de/OEKOBAU.DAT/resource/'
             'processes/<UUID>?format=json&view=extended '
             '– Indikator: GWP-total, Modul A1-A3')
ws4['A2'].font = Font(italic=True, size=9)
cols4 = ['Materialgruppe', 'Datensatz (Name)', 'UUID', 'Bezugseinheit',
         'Menge pro Datensatz', 'GWP-total A1-A3 (kg CO₂e/Datensatz)',
         'GWP pro Bezugseinheit', 'Verwendung']
for j, h in enumerate(cols4, 1):
    ws4.cell(row=4, column=j, value=h)
style_header(ws4, 4, len(cols4))
rows4 = [
    ['Baustahl', proc['Baustahl_Profile_Stabstahl']['name'],
     proc['Baustahl_Profile_Stabstahl']['uuid'], 'kg', 1000.0,
     proc['Baustahl_Profile_Stabstahl']['gwp']['A1-A3'],
     proc['Baustahl_Profile_Stabstahl']['gwp_per_unit'],
     'Standard – Flachstäbe, RO-Rundrohre, RD-12-Zugstäbe'],
    ['Baustahl (Alternative)', proc['Baustahl_Grobblech']['name'],
     proc['Baustahl_Grobblech']['uuid'], 'kg', 1000.0,
     proc['Baustahl_Grobblech']['gwp']['A1-A3'],
     proc['Baustahl_Grobblech']['gwp_per_unit'],
     'Alternative – Baustahl: Grobblech'],
    ['Beton', proc['Beton_C2530']['name'],
     proc['Beton_C2530']['uuid'], 'm³', 1.0,
     proc['Beton_C2530']['gwp']['A1-A3'],
     proc['Beton_C2530']['gwp_per_unit'],
     'Standard – Fundamente (Rectangle 500/500)'],
    ['Beton (Alternative)', proc['Beton_EcoPact']['name'],
     proc['Beton_EcoPact']['uuid'], 'm³', 1.0,
     proc['Beton_EcoPact']['gwp']['A1-A3'],
     proc['Beton_EcoPact']['gwp_per_unit'],
     'Alternative – Beton C25/30 EcoPact (Herstellerdatensatz)'],
    ['Betonstahl', proc['Betonstahl_Staeber']['name'],
     proc['Betonstahl_Staeber']['uuid'], 'kg', 1.0,
     proc['Betonstahl_Staeber']['gwp']['A1-A3'],
     proc['Betonstahl_Staeber']['gwp_per_unit'],
     'Standard – Bewehrung der Fundamente (pauschal 85 kg/m³ Beton)'],
]
for i, vals in enumerate(rows4, 1):
    for j, v in enumerate(vals, 1):
        cell = ws4.cell(row=4 + i, column=j, value=v)
        cell.border = thin
        cell.alignment = Alignment(wrap_text=True, vertical='top')
        if j in (5, 6, 7):
            cell.number_format = '#,##0.000000' if j == 7 else '#,##0.000'
for j, w in enumerate([20, 44, 38, 12, 14, 18, 16, 46], 1):
    ws4.column_dimensions[get_column_letter(j)].width = w

wb.save('Abgabe_23.10/EmbodiedCarbon_Staebe_Liste.xlsx')
print('Gespeichert: Abgabe_23.10/EmbodiedCarbon_Staebe_Liste.xlsx')
